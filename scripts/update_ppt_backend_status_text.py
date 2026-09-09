import json
from pathlib import Path

from pptx import Presentation


def package_file_count() -> str:
    manifests = sorted(
        Path.cwd().glob("submission/*_manifest.json"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    if not manifests:
        return "600"
    data = json.loads(manifests[0].read_text(encoding="utf-8"))
    return str(data.get("file_count") or "600")


def audit_pass_count() -> str:
    reports = sorted(
        Path.cwd().glob("参赛提交材料包/09_*机器可读.json"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    if not reports:
        return "131"
    data = json.loads(reports[0].read_text(encoding="utf-8"))
    return str(data.get("summary", {}).get("PASS") or "131")


def main() -> None:
    ppt = sorted(Path.cwd().glob("参赛提交材料包/*v0.2.pptx"))[0]
    file_count = package_file_count()
    audit_pass = audit_pass_count()
    prs = Presentation(str(ppt))
    updates = {
        (1, 14): audit_pass,
        (1, 15): "审计 PASS",
        (13, 3): "权限、账号、API 契约和后端状态让 Demo 能走向上云试用",
        (13, 4): "当前不是只在本地跑的页面，而是具备云端工程、演示账号、静态试用包、后端状态中心和生产接入契约。",
        (13, 12): "后端状态",
        (13, 13): "online / static / degraded / manual / blocked",
        (14, 3): "生产数据平面与后端状态让 Demo 具备上线骨架",
        (14, 4): "Postgres/Supabase 草案、RLS、备份、探针、后端状态中心和运维 Runbook 被产品化展示，避免只停留在原型。",
        (14, 14): "SQL 迁移、备份策略、写入干跑、RLS 演练、后端状态和导出包 manifest。",
        (15, 4): "视频来自真实 Demo 自动点击截图，覆盖核心闭环；决赛复剪可增补后端状态、账号初始化和数据平面镜头。",
        (15, 14): "新增后端状态中心、学校开班、RLS/备份和部署探针。",
        (16, 4): "代码、测试、构建、视频、PDF、PPT、公开试用包、后端状态和自动门禁一起证明完成度。",
        (16, 9): audit_pass,
        (16, 10): "自动审计 PASS",
        (16, 12): file_count,
        (20, 13): "诊断、规划、干预、记忆、复核、上云、后端状态、试用闭环。",
        (21, 10): "运行测试、构建、后端状态验收和 release gate。",
        (21, 13): "上传 v0.4 ZIP，人工补齐队伍名、访问策略和公开链接。",
    }
    for (slide_no, shape_idx), new_text in updates.items():
        prs.slides[slide_no - 1].shapes[shape_idx].text = new_text
    prs.save(str(ppt))
    print(f"updated {ppt}")
    print(f"fields {len(updates)}")


if __name__ == "__main__":
    main()

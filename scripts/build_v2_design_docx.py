from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"D:\相关比赛、论文\2026.8-自适应习路径决策与伴学")
DOC_DIR = ROOT / "SE-Path学伴_软件工程自适应学习伙伴智能体_产品设计方案"
MD_PATH = DOC_DIR / "SE-Path学伴_产品设计方案_v2_真实闭环参赛版.md"
DOCX_PATH = DOC_DIR / "SE-Path学伴_产品设计方案_v2_真实闭环参赛版.docx"
QA_PATH = DOC_DIR / "qa" / "v2_design_docx_metrics.json"


def set_run_font(run, size: int | None = None, bold: bool | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text.strip())
    set_run_font(run, 9, bold)


def is_table_start(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    current = lines[index].strip()
    nxt = lines[index + 1].strip()
    return current.startswith("|") and current.endswith("|") and re.fullmatch(r"\|[\s:\-|\u2014]+\|", nxt) is not None


def parse_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = index
    while i < len(lines):
        line = lines[i].strip()
        if not (line.startswith("|") and line.endswith("|")):
            break
        if not re.fullmatch(r"\|[\s:\-|\u2014]+\|", line):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            rows.append(cells)
        i += 1
    return rows, i


def add_paragraph(doc: Document, text: str, style: str | None = None, bold: bool = False) -> None:
    paragraph = doc.add_paragraph(style=style)
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_after = Pt(4)
    paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_run_font(run, 10 if style is None else None, bold)


def build_docx() -> dict[str, int]:
    text = MD_PATH.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()

    section = doc.sections[0]
    section.page_width = 11906
    section.page_height = 16838
    section.top_margin = 900
    section.bottom_margin = 900
    section.left_margin = 900
    section.right_margin = 900

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    for name, size in [("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True

    title_added = False
    heading_count = 0
    table_count = 0
    paragraph_count = 0
    i = 0

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        if not line or line == "---":
            i += 1
            continue

        if is_table_start(lines, i):
            rows, next_i = parse_table(lines, i)
            if rows:
                max_cols = max(len(row) for row in rows)
                table = doc.add_table(rows=len(rows), cols=max_cols)
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c in range(max_cols):
                        set_cell_text(table.cell(r, c), row[c] if c < len(row) else "", bold=(r == 0))
                table_count += 1
            i = next_i
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading_match:
            level = len(heading_match.group(1))
            content = heading_match.group(2).strip()
            if level == 1 and not title_added:
                paragraph = doc.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(content)
                set_run_font(run, 22, True)
                title_added = True
            else:
                if level == 1:
                    doc.add_section(WD_SECTION.NEW_PAGE)
                style = f"Heading {min(level, 3)}"
                add_paragraph(doc, content, style=style, bold=True)
            heading_count += 1
            i += 1
            continue

        bullet_match = re.match(r"^(\d+\.|-)\s+(.+)$", line)
        if bullet_match:
            add_paragraph(doc, bullet_match.group(2), style="List Bullet" if bullet_match.group(1) == "-" else "List Number")
            paragraph_count += 1
            i += 1
            continue

        add_paragraph(doc, line)
        paragraph_count += 1
        i += 1

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("SE-Path 学伴 v2 真实闭环参赛版")
    set_run_font(run, 9)

    doc.save(DOCX_PATH)

    with zipfile.ZipFile(DOCX_PATH) as zf:
        names = set(zf.namelist())
        required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
        missing = sorted(required - names)
        if missing:
            raise RuntimeError(f"Invalid DOCX package, missing: {missing}")

    metrics = {
        "markdown_chars": len(text),
        "headings": heading_count,
        "tables": table_count,
        "paragraphs": paragraph_count,
        "docx_bytes": DOCX_PATH.stat().st_size,
    }
    QA_PATH.parent.mkdir(parents=True, exist_ok=True)
    QA_PATH.write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8")
    return metrics


if __name__ == "__main__":
    try:
        result = build_docx()
    except Exception as exc:
        print(f"failed: {exc}", file=sys.stderr)
        raise
    print(json.dumps(result, ensure_ascii=False, indent=2))
